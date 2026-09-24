from __future__ import annotations

import tempfile
import unittest
import zipfile
import shutil
from pathlib import Path

from docx import Document

from sales_engineer.ai import GeminiClient
from sales_engineer.catalog import PRODUCTS
from sales_engineer.demo_case_nexuspay_v2 import build_nexuspay_v2
from sales_engineer.document_validation import extract_docx_text, normalize_document_text, validate_state_against_docx
from sales_engineer.institutional_policy import validate_institutional_whitelist
from sales_engineer.manifest import build_document_manifest
from sales_engineer.neutral_template_builder import _set_rows, _set_rows_dynamic, build_neutral_template_docx
from sales_engineer.pipeline import Pipeline, new_proposal


ROOT = Path(__file__).resolve().parents[1]


class PipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.proposal = new_proposal({
            "client_name": "Cliente Teste S.A.",
            "opportunity": "Proteção Akamai",
            "products": ["edge_dns", "app_api_protector"],
            "transcript": "O cliente precisa proteger APIs e manter o DNS disponível. Eventos devem seguir para o SIEM.",
            "requirements": "Proteger APIs; manter DNS disponível; integrar ao SIEM",
            "scope": "Configuração e homologação",
        })

    def test_full_pipeline_produces_traceability_and_sections(self) -> None:
        result = Pipeline(GeminiClient()).run(self.proposal)
        self.assertEqual(result["progress"], 100)
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["document"]["status"], "blocked")
        self.assertGreaterEqual(len(result["traceability"]), 3)
        self.assertGreaterEqual(len(result["sections"]), 6)
        self.assertGreater(result["quality_score"], 0)

    def test_docx_is_valid_and_contains_client(self) -> None:
        result = Pipeline(GeminiClient()).run(self.proposal)
        result["governance"]["warranty"] = {
            "state": "commitment",
            "value": "A POPULOS corrigirá defeitos diretamente atribuíveis aos serviços executados por 60 dias corridos após o aceite final.",
            "source": "aprovação comercial do caso de teste",
        }
        result["governance"]["license_supply"] = {
            "state": "commitment",
            "value": "As licenças Akamai serão fornecidas pela POPULOS por meio de revenda autorizada.",
            "source": "aprovação comercial do caso de teste",
        }
        template = ROOT / "assets" / "akamai_template.docx"
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "proposal.docx"
            manifest = build_document_manifest(result)
            build_neutral_template_docx(result, ROOT / "assets" / "proposal_neutral_template.docx", output)
            self.assertTrue(output.exists())
            self.assertGreater(output.stat().st_size, 20_000)
            with zipfile.ZipFile(output) as package:
                self.assertIsNone(package.testzip())
                document = package.read("word/document.xml").decode("utf-8")
            self.assertIn("Cliente Teste S.A.", document)
            self.assertNotIn("CPFL", document)
            self.assertNotIn("Citrix", document)
            self.assertIn("Dimensionamento e Abrangência", document)
            self.assertEqual(len(Document(output).sections), 2)
            self.assertIn("Resumo executivo", " ".join(section["title"] for section in result["sections"]))

    def test_manifest_contains_only_selected_products(self) -> None:
        result = Pipeline(GeminiClient()).run(self.proposal)
        manifest = build_document_manifest(result)
        solution = next(block for block in manifest["blocks"] if block["id"] == "solution")
        self.assertEqual(solution["products"], ["edge_dns", "app_api_protector"])
        self.assertNotIn("ip_accelerator", solution["products"])

    def test_portfolio_contains_cross_layer_controls(self) -> None:
        self.assertTrue({"gtm", "ion", "prolexic", "account_protector"}.issubset(PRODUCTS))

    def test_competitive_demo_has_traceability_and_production_phase(self) -> None:
        competitive = build_nexuspay_v2(self.proposal)
        statuses = {item["product_id"]: item["status"] for item in competitive["solution_decisions"]}
        self.assertEqual(statuses["gtm"], "recommended")
        self.assertEqual(statuses["prolexic"], "recommended")
        self.assertEqual(statuses["account_protector"], "optional")
        phase_names = [item["name"] for item in competitive["delivery"]["phases"]]
        self.assertIn("Onda 1 — pré-evento", phase_names)
        self.assertIn("Onda 2 — pós-evento", phase_names)
        self.assertEqual(len(competitive["traceability"]), 7)
        self.assertTrue(competitive["event_readiness"])
        self.assertEqual(competitive["governance"]["warranty"]["state"], "populos_internal_decision")

    def test_competitive_demo_blocks_docx_until_internal_decisions(self) -> None:
        competitive = build_nexuspay_v2(self.proposal)
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "decisão interna POPULOS"):
                build_neutral_template_docx(
                    competitive,
                    ROOT / "assets" / "proposal_neutral_template.docx",
                    Path(directory) / "blocked.docx",
                )

    def test_fixed_slots_reject_silent_truncation(self) -> None:
        document = Document()
        paragraphs = [document.add_paragraph(f"slot {index}") for index in range(3)]
        with self.assertRaisesRegex(ValueError, "Capacidade fixa excedida"):
            _set_rows(paragraphs, range(2), ["um", "dois", "três"])

    def test_dynamic_slots_preserve_scope_deliverables(self) -> None:
        document = Document()
        paragraphs = [document.add_paragraph(f"slot {index}") for index in range(4)]
        _set_rows_dynamic(paragraphs, range(2), ["um", "dois", "três"], paragraphs[2])
        self.assertIn("três", " ".join(item.text for item in document.paragraphs))

    def test_competitive_docx_keeps_all_deliverables(self) -> None:
        competitive = build_nexuspay_v2(self.proposal)
        competitive["governance"]["warranty"] = {
            "state": "commitment", "value": "Garantia de teste por 90 dias.", "source": "teste"
        }
        competitive["governance"]["license_supply"] = {
            "state": "commitment", "value": "Licenças fornecidas pela POPULOS.", "source": "teste"
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "competitive.docx"
            build_neutral_template_docx(
                competitive, ROOT / "assets" / "proposal_neutral_template.docx", output
            )
            document = Document(output)
            text = " ".join(item.text for item in document.paragraphs)
            complete_text = extract_docx_text(output)
            event_paragraphs = [item for item in document.paragraphs if item.text in competitive["event_readiness"]]
        self.assertIn("HLD, runbooks, matriz de rastreabilidade e plano de testes", text)
        self.assertIn("As-built, evidências de homologação e plano de reversão", text)
        self.assertIn("Global Traffic Management:", text)
        self.assertIn("Ion:", text)
        self.assertTrue(event_paragraphs)
        self.assertTrue(all(not item.text.startswith("•") for item in event_paragraphs))
        self.assertEqual(complete_text.count("Garantia de teste por 90 dias."), 1)

    def test_normalization_removes_renderer_noise(self) -> None:
        source = "Failover multi-\ncloud — em até 90 s."
        rendered = "1. Failover multicloud – em até 90 s."
        self.assertEqual(normalize_document_text(source), normalize_document_text(rendered))

    def test_state_docx_gate_detects_missing_client_visible_item(self) -> None:
        proposal = {
            "scope": {"included": ["Configuração de Edge DNS"], "deliverables": ["HLD final"]},
            "acceptance_criteria": [],
            "event_readiness": [],
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "missing.docx"
            document = Document()
            document.add_paragraph("Configuração de Edge DNS")
            document.save(output)
            with self.assertRaisesRegex(ValueError, "Divergência estado × DOCX"):
                validate_state_against_docx(proposal, output)

    def test_institutional_whitelist_matches_approved_template(self) -> None:
        manifest = validate_institutional_whitelist(ROOT / "assets" / "proposal_neutral_template.docx")
        self.assertEqual(manifest["template_version"], "2026.09.24-1")

    def test_institutional_whitelist_rejects_changed_template(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            changed = Path(directory) / "proposal_neutral_template.docx"
            shutil.copy2(ROOT / "assets" / "proposal_neutral_template.docx", changed)
            changed.write_bytes(changed.read_bytes() + b"changed")
            shutil.copy2(ROOT / "assets" / "institutional_whitelist.json", changed.with_name("institutional_whitelist.json"))
            with self.assertRaisesRegex(ValueError, "Versão do template não autorizada"):
                validate_institutional_whitelist(changed)


if __name__ == "__main__":
    unittest.main()
