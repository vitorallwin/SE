from __future__ import annotations

import json
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest import mock

from docx import Document

from sales_engineer import case_runner
from sales_engineer.case_runner import read_source_material, record_attempt, skill_bundle, validate_state
from sales_engineer.proposal_rules import validate_proposal_rules
from tests.test_v6_rules import resolved_v6

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def vertice_latest() -> dict:
    """A real authoring output (diverse-case round, case 01), upgraded to the latest contract."""
    return json.loads((FIXTURES / "vertice_v11.json").read_text(encoding="utf-8"))


def vertice_source() -> str:
    return (FIXTURES / "vertice_insumo.md").read_text(encoding="utf-8")


def resolved_v7() -> dict:
    """The closed NexusPay V6 state upgraded to the v7 contract (regression fixture only, never an authoring example)."""
    proposal = resolved_v6()
    proposal["rules_version"] = 7
    proposal["data_mode"] = "test"
    for key, decision in proposal["governance"].items():
        decision.update({
            "approved_by": "Responsável comercial", "approved_at": "2026-09-24", "mode": "test",
            "approval_record": "resposta explícita no chat de 24/09", "source": f"decisão do caso para {key}",
        })
    proposal["governance"]["sla"] = {"state": "commitment", "basis": "populos_standard", "value": "tabela institucional do template POPULOS",
                                     "source": "template POPULOS vigente", "template_version": "2026.09.24-1", "applies_to": ["phased"]}
    proposal["governance"]["phase1_estimate"].update({"approved_by": "Arquiteto responsável", "approver_role": "arquiteto"})
    proposal["estimates"][0].update({"owner": "Arquiteto responsável", "owner_role": "arquiteto"})
    proposal["optional_phase"]["waves"][0]["id"] = "onda-1"
    proposal["optional_phase"]["waves"][1]["id"] = "onda-2"
    for criterion in proposal["acceptance_criteria"]:
        if criterion["phase"] == "optional":
            criterion["option_id"] = "onda-2" if criterion["requirement_id"] in {"REQ-01", "REQ-04"} else "onda-1"
    track = proposal["fast_track"]
    track["id"] = "trilha-rapida"
    unknown_lead = {"licenciamento": {"weeks": None, "pending_question": "Qual o prazo de contrato e provisionamento das licenças?"}}
    track["lead_times"] = deepcopy(unknown_lead)
    track["items"].append("Licenciamento: contrato de revenda e provisionamento das licenças antes da janela")
    proposal["acceptance_criteria"].append({"requirement_id": "REQ-02", "phase": "optional", "option_id": "trilha-rapida",
                                            "criterion": "Hostnames prioritários na borda sem regressão funcional e reversão executada em ensaio."})
    proposal["event_feasibility"][0]["lead_times"] = deepcopy(unknown_lead)
    return proposal


class V7RulesTests(unittest.TestCase):
    def assertViolation(self, proposal: dict, fragment: str) -> None:
        errors = validate_proposal_rules(proposal)
        self.assertTrue(any(fragment in error for error in errors), f"esperado '{fragment}' em {errors}")

    def test_v7_fixture_is_clean(self) -> None:
        self.assertEqual(validate_proposal_rules(resolved_v7()), [])

    def test_v6_decisions_fail_v7(self) -> None:
        proposal = resolved_v6()
        proposal["rules_version"] = 7
        errors = " ".join(validate_proposal_rules(proposal))
        for fragment in ("data_mode ausente", "sem registro da aprovação explícita", "aprovação implícita", "estimativa sem dono técnico"):
            self.assertIn(fragment, errors)

    # Aprovação explícita
    def test_implicit_approval_blocks(self) -> None:
        proposal = resolved_v7()
        proposal["governance"]["engagement_type"]["approval_record"] = "aprovado ao seguir para a próxima versão"
        self.assertViolation(proposal, "aprovação implícita ou marcação de teste no texto (approval_record): engagement_type")

    def test_invalid_approval_date_blocks(self) -> None:
        proposal = resolved_v7()
        proposal["governance"]["warranty"]["approved_at"] = "ontem"
        self.assertViolation(proposal, "decisão sem data de aprovação válida (AAAA-MM-DD): warranty")

    # Modo de teste é flag, não texto, e não desliga checagem
    def test_test_marker_in_text_blocks(self) -> None:
        proposal = resolved_v7()
        proposal["governance"]["warranty"]["source"] = "dado de teste"
        self.assertViolation(proposal, "(source): warranty")

    # Padrão institucional: versão do template, não aprovação por caso
    def test_populos_standard_needs_current_template_version_not_case_approval(self) -> None:
        proposal = resolved_v7()
        self.assertNotIn("approved_by", proposal["governance"]["sla"])
        proposal["governance"]["sla"]["template_version"] = "2025.01.01-1"
        self.assertViolation(proposal, "padrão institucional sem a versão vigente do template (2026.09.24-1): sla")

    def test_technical_role_is_normalized(self) -> None:
        proposal = resolved_v7()
        proposal["estimates"][0]["owner_role"] = "Arquiteta de Soluções"
        proposal["governance"]["phase1_estimate"]["approver_role"] = "Engenharia"
        self.assertEqual(validate_proposal_rules(proposal), [])

    def test_test_mode_does_not_disable_checks(self) -> None:
        proposal = resolved_v7()
        proposal["traceability"][0]["source"] = ""
        self.assertEqual(proposal["data_mode"], "test")
        self.assertViolation(proposal, "requisito sem referência de origem: REQ-01")

    # Estimativa com dono técnico
    def test_effort_estimate_needs_technical_owner(self) -> None:
        proposal = resolved_v7()
        proposal["estimates"][0]["owner_role"] = "comercial"
        proposal["governance"]["phase1_estimate"]["approver_role"] = "comercial"
        self.assertViolation(proposal, "estimativa sem dono técnico (owner_role)")
        self.assertViolation(proposal, "estimativa de esforço aprovada por papel não técnico: phase1_estimate")

    # Opção que toca produção tem aceite próprio
    def test_production_option_needs_own_acceptance(self) -> None:
        proposal = resolved_v7()
        proposal["acceptance_criteria"] = [c for c in proposal["acceptance_criteria"] if c.get("option_id") != "trilha-rapida"]
        self.assertViolation(proposal, "opção que toca produção sem critério de aceite próprio: trilha-rapida")

    def test_optional_criterion_must_point_to_existing_option(self) -> None:
        proposal = resolved_v7()
        proposal["acceptance_criteria"][-1]["option_id"] = "inexistente"
        self.assertViolation(proposal, "critério de fase opcional sem opção existente: inexistente")

    # Licenciamento na viabilidade
    def test_license_lead_time_required(self) -> None:
        proposal = resolved_v7()
        proposal["event_feasibility"][0]["lead_times"] = {}
        proposal["fast_track"]["lead_times"] = {"licenciamento": {"weeks": None}}
        proposal["fast_track"]["items"] = [i for i in proposal["fast_track"]["items"] if "icen" not in i]
        self.assertViolation(proposal, "viabilidade sem prazo de licenciamento declarado")
        self.assertViolation(proposal, "trilha rápida com prazo de licenciamento desconhecido sem pergunta aberta")
        self.assertViolation(proposal, "trilha rápida não lista o licenciamento entre os pré-requisitos")

    def test_known_license_lead_time_needs_source(self) -> None:
        proposal = resolved_v7()
        proposal["fast_track"]["lead_times"] = {"licenciamento": {"weeks": 2}}
        self.assertViolation(proposal, "trilha rápida com prazo de licenciamento sem origem")

    def test_unknown_license_lead_time_cannot_be_fits(self) -> None:
        proposal = resolved_v7()
        proposal["event_feasibility"][0]["classification"] = "fits"
        self.assertViolation(proposal, "viabilidade 'fits' com prazo de licenciamento desconhecido")


class CaseRunnerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.run_dir = Path(self.tmp.name) / "run"

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _submit(self, state) -> dict:
        self.run_dir.mkdir(parents=True, exist_ok=True)
        (self.run_dir / "insumo.md").write_text(vertice_source(), encoding="utf-8")
        path = Path(self.tmp.name) / "state.json"
        path.write_text(state if isinstance(state, str) else json.dumps(state, ensure_ascii=False), encoding="utf-8")
        return record_attempt(self.run_dir, path)

    def test_clean_attempt_is_composed(self) -> None:
        result = self._submit(vertice_latest())
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["composition"]["status"], "emitted", result["composition"])
        text = " ".join(p.text for p in Document(result["composition"]["docx"]).paragraphs)
        self.assertIn("Vértice Educação", text)
        self.assertTrue((self.run_dir / "final" / "relatorio-cobertura.md").exists())

    def test_iteration_limit_is_enforced(self) -> None:
        with mock.patch.object(case_runner, "engine_config", return_value={"max_iterations": 2}):
            self.assertEqual(self._submit("{").get("status"), "violations")
            self.assertEqual(self._submit({"rules_version": 11}).get("status"), "violations")
            self.assertEqual(self._submit(vertice_latest())["status"], "limit_reached")
        log = json.loads((self.run_dir / "attempts.json").read_text(encoding="utf-8"))
        self.assertEqual([a["attempt"] for a in log["attempts"]], [1, 2])
        self.assertTrue((self.run_dir / "pass1.json").exists())

    def test_malformed_and_old_states_come_back_as_feedback(self) -> None:
        self.assertIn("não é um objeto JSON", validate_state([])[0])
        self.assertIn("rules_version deve ser 11", validate_state({"rules_version": 10})[0])
        self.assertEqual(validate_state({"rules_version": 11, "input_assessment": "x"}),
                         ["formato (input_assessment): 'x' is not of type 'object'"])

    def test_source_extraction_reads_text_and_docx(self) -> None:
        source = Path(self.tmp.name) / "insumo"
        source.mkdir()
        (source / "email.txt").write_text("Pedido do cliente", encoding="utf-8")
        document = Document()
        document.add_paragraph("Ata da reunião")
        document.save(source / "ata.docx")
        (source / "planilha.xlsx").write_bytes(b"x")
        text = read_source_material(source)
        self.assertIn("Pedido do cliente", text)
        self.assertIn("Ata da reunião", text)
        self.assertIn("formato não suportado pelo extrator: planilha.xlsx", text)

    def test_workspace_is_physically_isolated_and_self_sufficient(self) -> None:
        import subprocess, sys
        source = Path(self.tmp.name) / "caso"
        source.mkdir()
        (source / "material.md").write_text(vertice_source(), encoding="utf-8")
        workspace = Path(self.tmp.name) / "ws"
        run_dir = case_runner.make_workspace(workspace, source, "2026-09-24", "test", "B")
        files = {p.relative_to(workspace).as_posix() for p in workspace.rglob("*") if p.is_file()}
        for forbidden in ("tests/", "data/", "demo_case", ".git", "generate_nexuspay", ".env"):
            self.assertFalse([f for f in files if forbidden in f], forbidden)
        request = json.loads((run_dir / "request.json").read_text(encoding="utf-8"))
        self.assertEqual(request["template_version"], "2026.09.24-1")
        self.assertEqual(request["skill_sha256"], case_runner.skill_hash())
        state = Path(self.tmp.name) / "estado.json"
        state.write_text(json.dumps(vertice_latest(), ensure_ascii=False), encoding="utf-8")
        result = subprocess.run([sys.executable, "scripts/validate_attempt.py", "execucao", str(state)],
                                cwd=workspace, capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertTrue((run_dir / "final" / "relatorio-cobertura.md").exists())

    def test_sealed_workspace_has_no_engine_code_and_still_validates(self) -> None:
        import subprocess, sys
        source = Path(self.tmp.name) / "caso"
        source.mkdir()
        (source / "material.md").write_text(vertice_source(), encoding="utf-8")
        workspace = Path(self.tmp.name) / "ws"
        run_dir = case_runner.make_workspace(workspace, source, "2026-09-24", "test", "B", sealed=True)
        files = {p.relative_to(workspace).as_posix() for p in workspace.rglob("*") if p.is_file()}
        self.assertEqual([f for f in files if f.endswith(".py")], ["scripts/validate_attempt.py"])
        self.assertFalse([f for f in files if f.startswith(("sales_engineer/", "assets/", "engine_config"))])
        state = run_dir / "estado.json"
        state.write_text(json.dumps(vertice_latest(), ensure_ascii=False), encoding="utf-8")
        result = subprocess.run([sys.executable, "scripts/validate_attempt.py", "execucao", "execucao/estado.json"],
                                cwd=workspace, capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertTrue(list((run_dir / "final").glob("*.docx")))

    def test_author_bundle_carries_no_case_material(self) -> None:
        bundle = skill_bundle()
        self.assertIn("state-schema.md", bundle)
        for term in ("NexusPay", "Black Friday 2026", "PT-2609"):
            self.assertNotIn(term, bundle)


if __name__ == "__main__":
    unittest.main()
