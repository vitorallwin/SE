from __future__ import annotations

import json
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from docx import Document

from scripts.generate_nexuspay_v5 import CASE_DECISIONS, ESTIMATE_OWNER
from sales_engineer.demo_case_nexuspay_v5 import build_nexuspay_v5
from sales_engineer.document_validation import detect_skill_leakage, extract_docx_text, find_vocabulary_violations
from sales_engineer.neutral_template_builder import _diagram_boxes, build_neutral_template_docx
from sales_engineer.proposal_rules import validate_proposal_rules, week_ranges


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "assets" / "proposal_neutral_template.docx"
SOURCE = ROOT / "data" / "proposals" / "20260923-e54ccb6.json"


def resolved_v5() -> dict:
    proposal = build_nexuspay_v5(json.loads(SOURCE.read_text(encoding="utf-8")))
    proposal["governance"].update(deepcopy(CASE_DECISIONS))
    for estimate in proposal["estimates"]:
        estimate["owner"] = ESTIMATE_OWNER
    return proposal


class V5StateRulesTests(unittest.TestCase):
    def assertViolation(self, proposal: dict, fragment: str) -> None:
        errors = validate_proposal_rules(proposal)
        self.assertTrue(any(fragment in error for error in errors), f"esperado '{fragment}' em {errors}")

    def test_resolved_v5_state_has_no_violations(self) -> None:
        self.assertEqual(validate_proposal_rules(resolved_v5()), [])

    def test_unresolved_case_decisions_block(self) -> None:
        proposal = build_nexuspay_v5(json.loads(SOURCE.read_text(encoding="utf-8")))
        errors = validate_proposal_rules(proposal)
        for key in ("engagement_type", "phase1_estimate", "warranty", "license_supply"):
            self.assertIn(f"decisão interna aberta: {key}", errors)
        self.assertIn("estimativa sem responsável: Duração da Fase 1", errors)

    # F-01
    def test_committed_implementation_against_client_scope_blocks(self) -> None:
        proposal = resolved_v5()
        proposal["delivery"]["phases"].append({"name": "Onda 1 — pré-evento", "kind": "implementation", "weeks": [0, 0]})
        self.assertViolation(proposal, "fase contratada com implantação contra o escopo do cliente")

    def test_engagement_must_be_registered(self) -> None:
        proposal = resolved_v5()
        proposal["governance"]["engagement_type"]["value"] = "inferido"
        self.assertViolation(proposal, "engagement_type ausente ou inválido")

    # Q-01
    def test_week_ranges_must_match_sum_of_phases(self) -> None:
        proposal = resolved_v5()
        proposal["sections"][0]["paragraphs"][1] += " A implantação está estimada entre oito e doze semanas."
        self.assertViolation(proposal, "faixa 8–12 semanas diverge da soma das fases (4–6)")

    def test_week_range_parser_reads_words_and_digits(self) -> None:
        self.assertEqual(week_ranges("entre quatro e seis semanas; 2 a 3 sem."), [(4, 6), (2, 3)])

    # Q-02 / Q-03
    def test_recommended_product_must_be_in_exactly_one_wave(self) -> None:
        proposal = resolved_v5()
        proposal["optional_phase"]["waves"][1]["components"].remove("edge_dns")
        self.assertViolation(proposal, "produto recomendado em 0 ondas (esperado 1): edge_dns")
        proposal["optional_phase"]["waves"][0]["components"].append("gtm")
        self.assertViolation(proposal, "produto recomendado em 2 ondas (esperado 1): gtm")

    def test_production_wave_needs_acceptance(self) -> None:
        proposal = resolved_v5()
        proposal["optional_phase"]["waves"][0]["acceptance"] = ""
        self.assertViolation(proposal, "onda em produção sem marco de aceite")

    # Q-04
    def test_optional_product_cannot_enter_traceability(self) -> None:
        proposal = resolved_v5()
        proposal["traceability"][0]["solution"] = "Edge DNS + GTM; ALB opcional"
        self.assertViolation(proposal, "opcional na matriz de rastreabilidade")

    # Provenance, F-06, F-07, F-02
    def test_requirement_needs_source_reference(self) -> None:
        proposal = resolved_v5()
        proposal["traceability"][0]["source"] = ""
        self.assertViolation(proposal, "requisito sem referência de origem: REQ-01")

    def test_pain_number_must_appear_in_summary(self) -> None:
        proposal = resolved_v5()
        proposal["sections"][0]["paragraphs"][0] = "A NexusPay precisa de resiliência."
        self.assertViolation(proposal, "número de dor ausente do sumário")

    def test_estimated_dimensioning_without_owner_blocks(self) -> None:
        proposal = resolved_v5()
        proposal["dimensioning"].append(("Aplicações e APIs prioritárias", "5 a 15", "Faixa", "Estimativa"))
        self.assertViolation(proposal, "dimensionamento estimado sem premissa de responsável")

    def test_siem_export_requires_masking(self) -> None:
        proposal = resolved_v5()
        proposal["scope"]["included"] = [item.replace("com mascaramento de dados de cartão e PII antes do envio", "com campos acordados") for item in proposal["scope"]["included"]]
        self.assertViolation(proposal, "integração SIEM sem controle de dados sensíveis")

    def test_every_client_requirement_has_a_req(self) -> None:
        proposal = resolved_v5()
        refs = {entry["statement"]: entry for entry in proposal["source_coverage"]}
        for statement in ("Dados de cartão e PII não podem ficar expostos em logs", "Dados sensíveis não podem ser interceptados no trânsito", "Regulação do Banco Central exige auditoria rigorosa", "Automatizar a resiliência sem impactar o SLA"):
            self.assertEqual(refs[statement]["status"], "coberto")


class V5DocumentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.proposal = resolved_v5()
        cls.directory = tempfile.TemporaryDirectory()
        cls.output = Path(cls.directory.name) / "v5.docx"
        build_neutral_template_docx(cls.proposal, TEMPLATE, cls.output)
        cls.text = extract_docx_text(cls.output)
        cls.paragraphs = [p.text for p in Document(cls.output).paragraphs]

    @classmethod
    def tearDownClass(cls) -> None:
        cls.directory.cleanup()

    def _section(self, start: str, end: str) -> str:
        begin = next(i for i, text in enumerate(self.paragraphs) if text.startswith(start))
        finish = next(i for i, text in enumerate(self.paragraphs) if i > begin and text.startswith(end))
        return " ".join(self.paragraphs[begin:finish])

    def test_optional_products_stay_out_of_primary_fronts_and_diagram(self) -> None:
        fronts = self._section("3.1 ", "3.4 ")
        self.assertNotIn("Application Load Balancer", fronts)
        self.assertNotIn("ALB", fronts)
        self.assertNotIn("Account Protector", fronts)
        labels = " ".join(sub for _, sub in _diagram_boxes(self.proposal))
        self.assertNotIn("ALB", labels)
        self.assertNotIn("Account", labels)
        self.assertIn("Application Load Balancer — incluído somente se", self._section("3.4 Opções", "3.5 "))

    def test_every_recommended_product_is_named_in_the_optional_phase(self) -> None:
        phase2 = self._section("6.3 ", "7.")
        for name in ("Edge DNS", "Global Traffic Management", "Ion", "App & API Protector", "Prolexic", "Bot Manager Premier"):
            self.assertEqual(phase2.count(f"Componentes:") , 2)
            self.assertIn(name, phase2)
        self.assertNotIn("demais componentes", self.text.casefold())

    def test_committed_scope_has_no_production_change(self) -> None:
        scope = self._section("6.1 ", "6.2 ")
        for term in ("Onboarding", "Configuração de propriedades", "em produção antes do freeze"):
            self.assertNotIn(term, scope)

    def test_all_requirements_rendered(self) -> None:
        for index in range(1, 11):
            self.assertIn(f"REQ-{index:02d}", self.text)

    def test_values_without_origin_are_gone(self) -> None:
        for value in ("90 segundos", "cinco minutos", "3 a 8", "5 a 15", "oito e doze", "oito a doze"):
            self.assertNotIn(value, self.text)

    def test_client_pain_numbers_in_summary(self) -> None:
        summary = self._section("2.", "Objetivos")
        self.assertIn("dois segundos", summary)
        self.assertIn("minutos preciosos", summary)

    def test_language_rules(self) -> None:
        self.assertEqual(find_vocabulary_violations(self.text, strict=True), [])
        self.assertNotIn("OBJ-", self.text)
        self.assertEqual(detect_skill_leakage(self.text), [])

    def test_warranty_paragraph_and_box_do_not_duplicate(self) -> None:
        self.assertEqual(self.text.count("90 dias corridos"), 1)
        self.assertNotIn("aprovado para esta", self.text)

    def test_builder_blocks_rule_violation(self) -> None:
        broken = deepcopy(self.proposal)
        broken["traceability"][2]["source"] = ""
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "Regras v5 violadas"):
                build_neutral_template_docx(broken, TEMPLATE, Path(directory) / "x.docx")


class LanguageDetectorTests(unittest.TestCase):
    def test_strict_vocabulary_catches_untranslated_and_internal_terms(self) -> None:
        found = find_vocabulary_violations("Conduzir war room conforme o compromisso aprovado para esta oportunidade.", strict=True)
        self.assertIn("war room", found)
        self.assertIn("oportunidade", found)

    def test_skill_leakage_detects_eight_copied_words(self) -> None:
        copied = "O texto diz: Every requirement and every client_fact carries a source reference such as transcrição."
        self.assertTrue(detect_skill_leakage(copied))
        self.assertEqual(detect_skill_leakage("Texto original sem nenhuma cópia das regras internas de autoria."), [])


if __name__ == "__main__":
    unittest.main()
