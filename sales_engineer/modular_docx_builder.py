from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from .catalog import PRODUCTS


NAVY = "151C32"
TEAL = "15A6A6"
ORANGE = "F47B20"
CYAN = "45C4E8"
LIGHT = "F2F6F8"
GRID = "D9E1E5"
TEXT = "1D2433"


def _font(size: int, bold: bool = False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _wrap(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> list[str]:
    words = str(text).split()
    rows: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= width:
            current = candidate
        else:
            if current:
                rows.append(current)
            current = word
    if current:
        rows.append(current)
    return rows


def _cover_image(proposal: dict[str, Any], manifest: dict[str, Any]) -> BytesIO:
    width, height = 1240, 1754
    image = Image.new("RGB", (width, height), "#151C32")
    draw = ImageDraw.Draw(image)
    for y in range(height):
        blend = y / height
        color = (21 + int(4 * blend), 28 + int(5 * blend), 50 + int(14 * blend))
        draw.line((0, y, width, y), fill=color)

    draw.ellipse((105, 92, 161, 148), outline="#FFFFFF", width=3)
    draw.text((185, 94), "POPULOS", font=_font(31), fill="#FFFFFF")
    draw.multiline_text((890, 92), "FREEDOM TO\nINNOVATE AND\nWORK WITHOUT\nLIMITATIONS", font=_font(23), fill="#C7CBD4", spacing=9, align="right")

    draw.text((110, 635), "PROPOSTA TÉCNICA", font=_font(57), fill="#FFFFFF")
    title_font = _font(38, bold=True)
    y = 790
    for row in _wrap(draw, manifest.get("subtitle", "Solução Akamai"), title_font, 960):
        draw.text((110, y), row, font=title_font, fill="#FFFFFF")
        y += 50
    draw.text((110, y + 24), f"• {manifest.get('client', '')}", font=_font(31, bold=True), fill="#FFFFFF")

    month_names = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
    now = datetime.now()
    draw.text((110, y + 82), f"{month_names[now.month - 1]} {now.year}", font=_font(21), fill="#D8DBE2")
    draw.text((110, y + 124), f"{manifest.get('proposal_code', '')}  |  versão {manifest.get('version', '1.0')}", font=_font(18), fill="#9EA6B7")

    for index, color in enumerate(("#15A6A6", "#45C4E8", "#F47B20")):
        draw.arc((-260 + index * 30, 1240 + index * 22, 1480 - index * 15, 1990 - index * 10), 190, 345, fill=color, width=3)
    draw.line((1160, 260, 1160, 760), fill="#45C4E8", width=3)
    draw.ellipse((1154, 748, 1166, 760), fill="#F47B20")

    stream = BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color: str = GRID) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), color)


def _set_run_font(run, name: str = "Arial", size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 30, "000000"), ("Heading 1", 21, "000000"), ("Heading 2", 14, TEXT), ("Heading 3", 11, TEAL)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 4)
        style.paragraph_format.space_after = Pt(6)
    title_pr = doc.styles["Title"]._element.get_or_add_pPr()
    border = title_pr.find(qn("w:pBdr"))
    if border is not None:
        title_pr.remove(border)


def _accent(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [Cm(0.45), Cm(0.45), Cm(0.45)]
    for cell, width, color in zip(table.rows[0].cells, widths, (TEAL, CYAN, ORANGE)):
        cell.width = width
        _set_cell_shading(cell, color)
        _set_cell_margins(cell, 0, 0, 0, 0)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    _set_table_borders(table, "FFFFFF")


def _heading(doc: Document, title: str, level: int = 1) -> None:
    if level == 1:
        _accent(doc)
    paragraph = doc.add_heading(title, level=level)
    paragraph.paragraph_format.keep_with_next = True


def _paragraph(doc: Document, text: str, bold_lead: str | None = None) -> None:
    if not str(text).strip():
        return
    paragraph = doc.add_paragraph()
    if bold_lead:
        run = paragraph.add_run(bold_lead)
        _set_run_font(run, bold=True)
    run = paragraph.add_run(str(text).strip())
    _set_run_font(run)


def _bullets(doc: Document, items: Iterable[Any]) -> None:
    for item in items or []:
        text = str(item).strip()
        if not text:
            continue
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        _set_run_font(paragraph.add_run(text))


def _table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_shading(cell, NAVY)
        _set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(header)
        _set_run_font(run, size=8.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                _set_cell_shading(cell, LIGHT)
            run = cell.paragraphs[0].add_run(str(value or ""))
            _set_run_font(run, size=8.2)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)


def _header_footer(section, proposal: dict[str, Any]) -> None:
    section.header.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(header.add_run(f"POPULOS  |  {proposal.get('code', '')}"), size=8, bold=True, color="5E6878")

    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(footer.add_run("Confidencial  •  "), size=8, color="6D7583")
    _add_page_field(footer)


def _render_block(doc: Document, block: dict[str, Any], proposal: dict[str, Any]) -> None:
    _heading(doc, str(block.get("title", "Seção")))
    kind = block.get("type")
    if kind == "narrative":
        for paragraph in block.get("paragraphs", []):
            _paragraph(doc, paragraph)
    elif kind == "discovery":
        for label, key in (("Objetivos", "business_goals"), ("Pontos de atenção", "pain_points"), ("Restrições", "constraints")):
            values = block.get(key, [])
            if values:
                _heading(doc, label, 2)
                _bullets(doc, values)
    elif kind == "solution":
        _paragraph(doc, block.get("summary", ""))
        products = block.get("products", [])
        flow = ["Clientes"]
        if "edge_dns" in products:
            flow.append("Edge DNS")
        if "alb" in products:
            flow.append("Roteamento global")
        if any(item in products for item in ("app_api_protector", "bot_manager", "malware_protection")):
            flow.append("Segurança na borda")
        if "ip_accelerator" in products:
            flow.append("Otimização TCP UDP")
        flow.append("Origens do cliente")
        _heading(doc, "Fluxo lógico", 2)
        _table(doc, [str(item) for item in flow], [], [6.6 / max(1, len(flow))] * len(flow))
        product_rows = []
        for decision in block.get("decisions", []):
            product = PRODUCTS.get(str(decision.get("product_id")), {})
            capabilities = decision.get("capabilities") or product.get("capabilities", [])
            product_rows.append([
                product.get("name", str(decision.get("product_id", "Solução"))),
                decision.get("rationale", ""),
                ", ".join(capabilities),
            ])
        _heading(doc, "Componentes recomendados", 2)
        _table(doc, ["Componente", "Justificativa", "Capacidades no escopo"], product_rows, [1.4, 3.2, 2.0])
    elif kind == "options":
        rows = []
        for decision in block.get("decisions", []):
            product = PRODUCTS.get(str(decision.get("product_id")), {})
            rows.append([product.get("name", decision.get("product_id")), decision.get("rationale", ""), ", ".join(decision.get("requirement_ids", []))])
        _table(doc, ["Opção", "Condição de inclusão", "Requisitos"], rows, [1.4, 4.0, 1.2])
    elif kind == "traceability":
        rows = []
        for item in block.get("rows", []):
            rows.append([item.get("requirement_id", ""), item.get("requirement", ""), item.get("solution", ""), item.get("evidence", ""), item.get("status", "")])
        _table(doc, ["ID", "Requisito", "Solução", "Evidência", "Cobertura"], rows, [0.6, 2.2, 1.2, 2.0, 0.7])
    elif kind == "scope":
        for label, key in (("Incluído", "included"), ("Entregáveis", "deliverables"), ("Fora do escopo", "excluded")):
            _heading(doc, label, 2)
            _bullets(doc, block.get(key, []))
    elif kind == "delivery":
        rows = []
        for phase in block.get("phases", []):
            rows.append([phase.get("name", ""), phase.get("objective", phase.get("duration", "")), ", ".join(phase.get("outputs", [])), phase.get("exit_criteria", "")])
        if rows:
            _table(doc, ["Fase", "Objetivo", "Saídas", "Condição de saída"], rows, [1.2, 2.0, 2.0, 1.5])
        if block.get("responsibilities"):
            _heading(doc, "Responsabilidades", 2)
            responsibility_rows = [[item.get("party", ""), item.get("responsibility", "")] for item in block["responsibilities"]]
            _table(doc, ["Parte", "Responsabilidade"], responsibility_rows, [1.5, 5.2])
    elif kind == "risk":
        if block.get("assumptions"):
            _heading(doc, "Premissas", 2)
            _bullets(doc, block["assumptions"])
        if block.get("risks"):
            _heading(doc, "Riscos e mitigação", 2)
            rows = [[item.get("risk", ""), item.get("impact", ""), item.get("mitigation", "")] for item in block["risks"]]
            _table(doc, ["Risco", "Impacto", "Mitigação"], rows, [2.4, 1.6, 2.7])
        if block.get("open_questions"):
            _heading(doc, "Pendências para confirmação", 2)
            _bullets(doc, block["open_questions"])
    elif kind == "acceptance":
        rows = []
        for item in block.get("criteria", []):
            rows.append([item.get("requirement_id", ""), item.get("criterion", ""), item.get("evidence", ""), item.get("owner", "")])
        _table(doc, ["Requisito", "Critério", "Evidência", "Responsável"], rows, [0.8, 2.8, 2.2, 1.0])


def build_modular_proposal_docx(proposal: dict[str, Any], manifest: dict[str, Any], output: Path) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _configure_styles(doc)

    cover = doc.sections[0]
    cover.page_width = Cm(21)
    cover.page_height = Cm(29.7)
    cover.top_margin = cover.bottom_margin = cover.left_margin = cover.right_margin = Cm(0)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.add_run().add_picture(_cover_image(proposal, manifest), width=Cm(21), height=Cm(29.55))

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.page_width = Cm(21)
    body.page_height = Cm(29.7)
    body.top_margin = Cm(1.9)
    body.bottom_margin = Cm(1.8)
    body.left_margin = Cm(2.0)
    body.right_margin = Cm(2.0)
    body.header_distance = Cm(0.7)
    body.footer_distance = Cm(0.7)
    _header_footer(body, proposal)

    title = doc.add_paragraph(style="Title")
    _set_run_font(title.add_run("Visão da proposta"), size=28, bold=True, color="000000")
    _paragraph(doc, "Este documento apresenta a solução técnica recomendada a partir dos requisitos registrados no discovery, incluindo escopo, rastreabilidade, riscos e critérios de aceite.")
    _table(
        doc,
        ["Cliente", "Oportunidade", "Código", "Versão", "Status"],
        [[manifest.get("client", ""), manifest.get("subtitle", ""), manifest.get("proposal_code", ""), manifest.get("version", "1.0"), manifest.get("status", "Para revisão")]],
        [1.4, 2.5, 1.0, 0.6, 1.2],
    )
    _heading(doc, "Conteúdo desta proposta", 2)
    for index, block in enumerate(manifest.get("blocks", []), 1):
        paragraph = doc.add_paragraph()
        _set_run_font(paragraph.add_run(f"{index:02d}  "), bold=True, color=TEAL)
        _set_run_font(paragraph.add_run(str(block.get("title", ""))), bold=True)

    forced_breaks = {"solution", "traceability", "delivery", "acceptance"}
    for index, block in enumerate(manifest.get("blocks", [])):
        if index == 0 or block.get("id") in forced_breaks:
            doc.add_page_break()
        _render_block(doc, block, proposal)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.core_properties.title = str(manifest.get("subtitle", "Proposta técnica Akamai"))
    doc.core_properties.subject = f"Proposta técnica para {manifest.get('client', '')}"
    doc.core_properties.author = "POPULOS Sales Engineering"
    doc.core_properties.keywords = "Akamai, proposta técnica, arquitetura, rastreabilidade"
    doc.save(output)
    return output
