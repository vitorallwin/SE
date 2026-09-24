from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
import re
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from .catalog import PRODUCTS
from .document_validation import detect_skill_leakage, find_vocabulary_violations, validate_state_against_docx
from .institutional_policy import validate_institutional_whitelist
from .proposal_rules import assert_proposal_rules


def _text(value: Any) -> str:
    return " ".join(str(value or "").split())


def _replace_paragraph(paragraph, value: Any) -> None:
    text = _text(value)
    rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        rpr = deepcopy(paragraph.runs[0]._r.rPr)
    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def _remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def _replace_cell(cell, value: Any) -> None:
    paragraph = cell.paragraphs[0]
    _replace_paragraph(paragraph, value)
    for extra in cell.paragraphs[1:]:
        _remove_paragraph(extra)


def _replace_run_tokens(paragraph, replacements: dict[str, str]) -> None:
    for run in paragraph.runs:
        if not run.text:
            continue
        value = run.text
        for source, target in replacements.items():
            value = value.replace(source, target)
        value = re.sub(r"«([^»]+)»", r"\1", value)
        if value != run.text:
            run.text = value


def _remove_table(table) -> None:
    parent = table._element.getparent()
    if parent is not None:
        parent.remove(table._element)


def _remove_row(row) -> None:
    parent = row._tr.getparent()
    if parent is not None:
        parent.remove(row._tr)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _trim_table_columns(table, keep: int) -> None:
    grid = table._tbl.tblGrid
    for node in list(grid)[keep:]:
        grid.remove(node)
    for row in table.rows:
        for cell in list(row._tr.tc_lst)[keep:]:
            row._tr.remove(cell)


def _replace_callout(cell, title: str, body: str) -> None:
    first = cell.paragraphs[0]
    _replace_paragraph(first, title)
    if first.runs:
        first.runs[0].bold = True
        first.runs[0].font.color.rgb = RGBColor(31, 42, 94)
    first.paragraph_format.keep_with_next = True
    first.paragraph_format.keep_together = True
    for extra in list(cell.paragraphs[1:]):
        _remove_paragraph(extra)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(_text(body))
    run.font.name = "Lato"
    run.font.size = Pt(8.5)
    run.bold = False


def _format_added_table(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_shading(cell, "1F2A5E" if row_index == 0 else ("F3F4F6" if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Lato"
                    run.font.size = Pt(7.5)
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(31, 42, 94)
                    run.bold = row_index == 0


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def _governance_value(proposal: dict[str, Any], key: str) -> Any:
    decision = proposal.get("governance", {}).get(key, {})
    if not decision or (decision.get("state") == "populos_internal_decision" and not decision.get("value")):
        raise ValueError(f"Emissão bloqueada por decisão interna POPULOS: {key}")
    return decision.get("value")


def _doc_text(proposal: dict[str, Any], key: str, default: str) -> str:
    """Case-specific client text from the state, falling back to the generic wording."""
    return str(proposal.get("document_text", {}).get(key) or default)


def _strict(proposal: dict[str, Any]) -> bool:
    return int(proposal.get("rules_version", 0) or 0) >= 5


_DIAGRAM_LABELS = {
    "ion": "Ion", "app_api_protector": "AAP", "bot_manager": "Bot Manager",
    "account_protector": "Account", "edge_dns": "Edge DNS", "gtm": "GTM", "alb": "ALB",
}


def _diagram_boxes(proposal: dict[str, Any]) -> list[tuple[str, str]]:
    """Only recommended products enter the main diagram; options stay in their own section."""
    recommended = [d.get("product_id") for d in proposal.get("solution_decisions", []) if d.get("status") == "recommended"]
    edge = [_DIAGRAM_LABELS[p] for p in ("ion", "app_api_protector", "bot_manager", "account_protector") if p in recommended]
    traffic = [_DIAGRAM_LABELS[p] for p in ("edge_dns", "gtm", "alb") if p in recommended]
    return [
        ("Clientes e lojistas", "Jornadas web e API"),
        ("Borda Akamai", " · ".join(edge) or "Controles de borda"),
        ("Decisão de tráfego", " · ".join(traffic) or "DNS atual"),
        (_doc_text(proposal, "diagram_origins_title", "Origens do cliente"), _doc_text(proposal, "diagram_origins", "Cloud A · Cloud B · DC")),
    ]


def _architecture_diagram(proposal: dict[str, Any]) -> BytesIO:
    width, height = 1800, 560
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "C:/Windows/Fonts/arial.ttf"
    bold_path = "C:/Windows/Fonts/arialbd.ttf"
    try:
        font = ImageFont.truetype(font_path, 30)
        bold = ImageFont.truetype(bold_path, 32)
    except OSError:
        font = bold = ImageFont.load_default()
    boxes = _diagram_boxes(proposal)
    x_positions = [40, 480, 930, 1380]
    for index, ((title, sub), x) in enumerate(zip(boxes, x_positions)):
        draw.rounded_rectangle((x, 125, x + 360, 405), radius=24, fill="#F3F4F6", outline="#6B46C1", width=4)
        tb = draw.textbbox((0, 0), title, font=bold)
        draw.text((x + (360 - (tb[2] - tb[0])) / 2, 205), title, font=bold, fill="#1F2A5E")
        sb = draw.textbbox((0, 0), sub, font=font)
        draw.text((x + (360 - (sb[2] - sb[0])) / 2, 285), sub, font=font, fill="#3F4A6B")
        if index < len(boxes) - 1:
            draw.line((x + 370, 265, x_positions[index + 1] - 10, 265), fill="#6B46C1", width=7)
            draw.polygon([(x_positions[index + 1] - 10, 265), (x_positions[index + 1] - 34, 251), (x_positions[index + 1] - 34, 279)], fill="#6B46C1")
    if any(d.get("product_id") == "prolexic" and d.get("status") == "recommended" for d in proposal.get("solution_decisions", [])):
        draw.text((40, 470), "Prolexic protege a infraestrutura exposta quando os requisitos de rede e roteamento forem validados.", font=font, fill="#1F2A5E")
    stream = BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _set_rows(paragraphs: list, indexes: Iterable[int], values: Iterable[Any]) -> None:
    values = list(values)
    indexes = list(indexes)
    if len(values) > len(indexes):
        raise ValueError(
            f"Capacidade fixa excedida: {len(values)} item(ns) para {len(indexes)} posição(ões)"
        )
    for index, value in zip(indexes, values):
        _replace_paragraph(paragraphs[index], value)
    for index in indexes[len(values):]:
        _remove_paragraph(paragraphs[index])


def _set_rows_dynamic(paragraphs: list, indexes: Iterable[int], values: Iterable[Any], anchor) -> None:
    """Preenche os slots existentes e insere o excedente sem truncar conteúdo."""
    values = list(values)
    indexes = list(indexes)
    existing = values[:len(indexes)]
    _set_rows(paragraphs, indexes, existing)
    if not values[len(indexes):]:
        return
    source = paragraphs[indexes[-1]] if indexes else anchor
    style = source.style
    paragraph_properties = deepcopy(source._p.pPr) if source._p.pPr is not None else None
    for value in values[len(indexes):]:
        paragraph = anchor.insert_paragraph_before(_text(value))
        paragraph.style = style
        if paragraph_properties is not None:
            current = paragraph._p.pPr
            if current is not None:
                paragraph._p.remove(current)
            paragraph._p.insert(0, deepcopy(paragraph_properties))
        paragraph.paragraph_format.keep_together = True


def _insert_paragraph_like(anchor, source, value: Any, copy_run_format: bool = False):
    """Insert content with the template's real list/paragraph properties (and run formatting, for headings)."""
    paragraph = anchor.insert_paragraph_before(_text(value))
    paragraph.style = source.style
    if source._p.pPr is not None:
        current = paragraph._p.pPr
        if current is not None:
            paragraph._p.remove(current)
        paragraph._p.insert(0, deepcopy(source._p.pPr))
    if copy_run_format and source.runs and source.runs[0]._r.rPr is not None and paragraph.runs:
        paragraph.runs[0]._r.insert(0, deepcopy(source.runs[0]._r.rPr))
    paragraph.paragraph_format.keep_together = True
    return paragraph


def _product_name(product_id: str) -> str:
    return str(PRODUCTS.get(product_id, {}).get("name", product_id))


_FRONT_SENTENCES = {
    "edge_dns": "Edge DNS mantém a camada autoritativa distribuída.",
    "gtm": "Global Traffic Management decide o destino entre as origens com base em saúde e política.",
    "alb": "Application Load Balancer roteia requisições HTTP entre origens.",
    "ion": "Ion otimiza a entrega das jornadas web priorizadas.",
    "app_api_protector": "App & API Protector aplica WAF adaptativo, descoberta de APIs, proteção de dados sensíveis e mitigação DDoS de camada 7.",
    "prolexic": "Prolexic amplia a cobertura para ataques volumétricos de rede e exposição direta das origens.",
    "bot_manager": "Bot Manager Premier classifica e responde à automação em cadastro e login.",
    "account_protector": "Account Protector avalia o risco por identidade na criação, autenticação e atividade pós-login.",
}


def _fronts(proposal: dict[str, Any]) -> list[dict[str, Any]]:
    """Front text is composed only from recommended products, so options never read as part of the solution."""
    decisions = [d for d in proposal.get("solution_decisions", []) if d.get("status") == "recommended"]
    groups = [
        ("Continuidade e desempenho das jornadas digitais", ["edge_dns", "gtm", "alb", "ion"]),
        ("Proteção de aplicações, APIs e infraestrutura", ["app_api_protector", "prolexic"]),
        ("Proteção contra automação e abuso de contas", ["bot_manager", "account_protector"]),
    ]
    fronts: list[dict[str, Any]] = []
    for title, ids in groups:
        members = sorted((d for d in decisions if d.get("product_id") in ids), key=lambda d: ids.index(d.get("product_id")))
        if members:
            summary = " ".join(d.get("client_summary") or _FRONT_SENTENCES.get(str(d.get("product_id")), "") for d in members)
            fronts.append({
                "title": title,
                "products": [_product_name(str(d.get("product_id"))) for d in members],
                "summary": summary,
                "items": [
                    f"{_product_name(str(d.get('product_id')))}: {', '.join(str(cap) for cap in d.get('capabilities', [])[:3])}"
                    for d in members
                ],
            })
    return fronts


def _client_mark(client: str) -> bytes:
    canvas = Image.new("RGB", (900, 900), "#F7F7FC")
    draw = ImageDraw.Draw(canvas)
    fonts = ["C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/calibrib.ttf"]
    font = None
    for path in fonts:
        try:
            font = ImageFont.truetype(path, 78)
            break
        except OSError:
            pass
    font = font or ImageFont.load_default()
    words = client.upper().split()
    lines = [" ".join(words[: max(1, len(words) // 2)]), " ".join(words[max(1, len(words) // 2):])]
    lines = [line for line in lines if line]
    heights = []
    for line in lines:
        box = draw.textbbox((0, 0), line, font=font)
        heights.append(box[3] - box[1])
    y = (900 - sum(heights) - 35 * (len(lines) - 1)) / 2
    for line, height in zip(lines, heights):
        box = draw.textbbox((0, 0), line, font=font)
        x = (900 - (box[2] - box[0])) / 2
        draw.text((x, y), line, font=font, fill="#202A5C")
        y += height + 35
    stream = BytesIO()
    canvas.save(stream, format="JPEG", quality=94)
    return stream.getvalue()


def _replace_client_mark(docx_path: Path, client: str) -> None:
    source = docx_path.read_bytes()
    inbound = BytesIO(source)
    outbound = BytesIO()
    with zipfile.ZipFile(inbound, "r") as zin, zipfile.ZipFile(outbound, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/media/image2.jpeg":
                data = _client_mark(client)
            zout.writestr(item, data)
    docx_path.write_bytes(outbound.getvalue())


def build_neutral_template_docx(proposal: dict[str, Any], template: Path, output: Path) -> Path:
    """Populate the approved POPULOS neutral template without using it as a product checklist."""
    if not template.exists():
        raise FileNotFoundError(f"Template neutro não encontrado: {template}")
    warranty = _governance_value(proposal, "warranty")
    license_supply = _governance_value(proposal, "license_supply")
    if _strict(proposal):
        assert_proposal_rules(proposal)
    validate_institutional_whitelist(template)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(template)
    p = list(doc.paragraphs)
    t = list(doc.tables)

    def heading2(anchor, text):
        return _insert_paragraph_like(anchor, p[85], text, copy_run_format=True)

    def heading3(anchor, text):
        return _insert_paragraph_like(anchor, p[92], text, copy_run_format=True)

    def body_like(anchor, text):
        return _insert_paragraph_like(anchor, p[93], text, copy_run_format=True)

    client = _text(proposal.get("client_name"))
    opportunity = _text(proposal.get("opportunity"))
    requirements = proposal.get("discovery", {}).get("requirements", [])
    goals = proposal.get("discovery", {}).get("business_goals", [])
    fronts = _fronts(proposal)
    decisions = [d for d in proposal.get("solution_decisions", []) if d.get("status") == "recommended"]
    product_names = [_product_name(str(d.get("product_id"))) for d in decisions]
    summary = next((s.get("paragraphs", []) for s in proposal.get("sections", []) if s.get("title", "").casefold() == "resumo executivo"), [])
    month = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"][datetime.now().month - 1]

    _replace_paragraph(p[11], f"{opportunity} · {client} · {proposal.get('code', '')}")
    _replace_paragraph(p[12], f"EMPRESA: {client}")
    _replace_paragraph(p[13], f"| {month.upper()} {datetime.now().year} |")

    _replace_paragraph(p[19], "Sobre a Akamai")
    _replace_paragraph(p[20], _doc_text(proposal, "about_akamai", f"A Akamai oferece serviços distribuídos de entrega, desempenho, proteção e observabilidade para aplicações, APIs e infraestrutura exposta à internet. A arquitetura desta proposta combina apenas as capacidades ligadas às jornadas e riscos priorizados por {client}."))
    _replace_paragraph(p[21], _doc_text(proposal, "about_solution", f"Para {client}, a solução integra continuidade multi-cloud, proteção de aplicações e APIs, defesa contra abuso automatizado e otimização das jornadas de checkout e cadastro. Cada frente possui escopo de configuração, evidência de teste e responsável definidos."))
    _replace_paragraph(p[22], "Parceria POPULOS e Akamai")
    _replace_paragraph(p[23], _doc_text(proposal, "partnership", "A POPULOS responde pela arquitetura, implantação, testes e documentação do projeto, com profissionais qualificados nas tecnologias previstas. As credenciais nominais da equipe serão apresentadas na mobilização, conforme os requisitos formais da contratação."))
    _set_rows(p, range(24, 28), ["Arquitetura e segurança de aplicações Akamai;", "DNS, gestão global de tráfego e entrega de aplicações;", "Gestão de projeto, testes e documentação técnica."])
    _replace_paragraph(p[28], "A composição final da equipe será confirmada no plano de projeto, preservando as competências exigidas para cada frente.")

    if summary:
        _replace_paragraph(p[31], summary[0])
        _replace_paragraph(p[32], summary[1] if len(summary) > 1 else "A solução proposta atende às seguintes frentes:")
        for extra in summary[2:]:
            _insert_paragraph_like(p[33], p[32], extra, copy_run_format=True)
    else:
        _replace_paragraph(p[31], f"Esta proposta técnica apresenta a abordagem recomendada para {client}.")
        _replace_paragraph(p[32], "O escopo foi construído a partir dos requisitos e jornadas críticas apresentados pela NexusPay.")
    front_lines = [f"{front['title']} — {', '.join(front['products'])}." for front in fronts]
    p[30].paragraph_format.keep_with_next = True
    _set_rows(p, range(33, 38), front_lines)
    _replace_paragraph(p[38], _doc_text(proposal, "coverage", "Abrangência: checkout, APIs de pagamento, cadastro e autenticação de lojistas, zonas DNS e origens distribuídas entre duas nuvens públicas e o datacenter corporativo."))
    _set_rows_dynamic(p, range(41, 48), goals, p[48])

    _replace_paragraph(p[51], proposal.get("architecture", {}).get("summary", ""))
    diagram = p[52].insert_paragraph_before()
    diagram.alignment = WD_ALIGN_PARAGRAPH.CENTER
    diagram.add_run().add_picture(_architecture_diagram(proposal), width=Inches(6.65))
    front_slots = [(52, 53, range(54, 57)), (57, 58, []), (59, 60, [])]
    for slot, front in zip(front_slots, fronts):
        hi, body, bullet_indexes = slot
        _replace_paragraph(p[hi], f"3.{fronts.index(front)+1} {front['title']}")
        _replace_paragraph(p[body], front["summary"])
        if bullet_indexes:
            _set_rows(p, bullet_indexes, front.get("items", []))
    for hi, body, bullet_indexes in front_slots[len(fronts):]:
        _remove_paragraph(p[hi]); _remove_paragraph(p[body])
        for ix in bullet_indexes: _remove_paragraph(p[ix])

    _replace_paragraph(p[61], "4. Dimensionamento e Abrangência")
    _replace_paragraph(p[62], _doc_text(proposal, "dimensioning_intro", "Para esta versão, o planejamento considera as jornadas críticas já identificadas e uma implantação em ondas. A quantidade contratual de zonas, propriedades, APIs, usuários e capacidade de proteção será fechada no levantamento inicial, sem interromper o desenho e a preparação técnica."))
    _replace_paragraph(p[63], "4.1 Itens de dimensionamento")
    _replace_paragraph(p[65], "Observação: a linha de base será levantada e validada durante o Assessment.")
    _replace_paragraph(p[66], "4.2 Componentes centralizados e serviços de borda")
    p[66].paragraph_format.page_break_before = False
    dims = proposal.get("dimensioning", [
        ("Zonas DNS e domínios críticos", "Faixa inicial: 3 a 8", "Validar inventário", "Premissa"),
        ("Aplicações e APIs prioritárias", "Faixa inicial: 5 a 15", "Validar catálogo", "Premissa"),
        ("Origens multi-cloud", "3 ambientes", "Duas nuvens e um datacenter", "Confirmado"),
    ])
    if len(dims) > len(t[3].rows) - 1:
        raise ValueError(
            f"Dimensionamento excede a capacidade do template: {len(dims)} itens para {len(t[3].rows) - 1} linhas"
        )
    _trim_table_columns(t[3], 4)
    for c, value in zip(t[3].rows[0].cells, ["Item", "Quantidade de referência", "Base de planejamento", "Classificação"]): _replace_cell(c, value)
    for row, values in zip(t[3].rows[1:], dims):
        for c, value in zip(row.cells, values): _replace_cell(c, value)
    component_rows = []
    for decision in proposal.get("solution_decisions", []):
        if decision.get("status") not in {"recommended", "optional"}:
            continue
        status = "Recomendado" if decision.get("status") == "recommended" else "Opção condicionada"
        if decision.get("implementation_status") == "modality_pending":
            status += " — modalidade de implantação a definir"
        component_rows.append((_product_name(str(decision.get("product_id"))), status))
    component_table = doc.add_table(rows=1, cols=2)
    p[67]._p.addprevious(component_table._tbl)
    for c, value in zip(component_table.rows[0].cells, ["Componente", "Posição na proposta"]): _replace_cell(c, value)
    for name, status in component_rows:
        cells = component_table.add_row().cells
        _replace_cell(cells[0], name); _replace_cell(cells[1], status)
    _format_added_table(component_table, [4.6, 2.0])
    _repeat_table_header(component_table.rows[0])
    _remove_table(t[4])

    optional_decisions = [d for d in proposal.get("solution_decisions", []) if d.get("status") == "optional"]
    readiness = proposal.get("event_readiness", [])
    next_number = iter(range(len(fronts) + 1, len(fronts) + 4))
    options_number = next(next_number) if optional_decisions else None
    readiness_number = next(next_number) if readiness else None
    trace_number = next(next_number)
    trace_heading = heading2(p[61], f"3.{trace_number} Matriz de rastreabilidade")
    trace_table = doc.add_table(rows=1, cols=4)
    trace_heading._p.addnext(trace_table._tbl)
    trace_headers = ["Requisito", "Necessidade", "Componente", "Evidência de aceite"]
    for c, value in zip(trace_table.rows[0].cells, trace_headers): _replace_cell(c, value)
    for item in proposal.get("traceability", []):
        row = trace_table.add_row().cells
        values = [item.get("requirement_id", ""), item.get("requirement", ""), item.get("solution", ""), item.get("evidence", "")]
        for c, value in zip(row, values): _replace_cell(c, value)
    _format_added_table(trace_table, [0.65, 2.2, 1.35, 2.4])
    _repeat_table_header(trace_table.rows[0])

    _replace_paragraph(p[68], "Como primeira etapa, a POPULOS realizará o levantamento detalhado do ambiente atual, contemplando no mínimo:")
    _set_rows(p, range(69, 72), proposal.get("assessment_items") or [
        "Inventário de domínios, aplicações, APIs e origens;",
        "Métricas de tráfego, latência, disponibilidade e eventos de segurança;",
        "Dependências de integração, acesso, observabilidade e conformidade.",
    ])
    included = proposal.get("scope", {}).get("included", [])
    if proposal.get("document_text", {}).get("heading_6"):
        _replace_paragraph(p[73], proposal["document_text"]["heading_6"])
    _replace_paragraph(p[74], _doc_text(proposal, "methodology_intro", "A POPULOS disponibilizará equipe técnica e gestão de projeto para conduzir Assessment, Design, Implementação, Homologação, Produção e Estabilização. A execução seguirá o plano aprovado, com controles de mudança e evidências de cada marco."))
    _set_rows_dynamic(p, range(76, 85), included + proposal.get("scope", {}).get("deliverables", []), p[85])
    phases = proposal.get("delivery", {}).get("phases", [])
    optional_phase = proposal.get("optional_phase") or {}
    if optional_phase.get("waves"):
        anchor = p[94]
        heading = heading2(anchor, f"6.3 {optional_phase.get('title', 'Fase opcional')}")
        heading.paragraph_format.keep_with_next = True
        for text in [optional_phase.get("intro", "")] + list(optional_phase.get("conditions", [])):
            if text:
                body_like(anchor, text)
        for wave in optional_phase["waves"]:
            wave_heading = heading3(anchor, str(wave.get("name", "Onda")))
            wave_heading.paragraph_format.keep_with_next = True
            components = ", ".join(_product_name(str(item)) for item in wave.get("components", []))
            for text in (f"Componentes: {components}.", wave.get("activities", ""), wave.get("acceptance", "")):
                if text:
                    body_like(anchor, text)
    fast_track = proposal.get("fast_track") or {}
    if fast_track:
        heading2(p[94], f"6.{4 if optional_phase.get('waves') else 3} {fast_track.get('title', 'Trilha rápida opcional')}")
        for text in fast_track.get("paragraphs", []):
            body_like(p[94], text)
        for item in fast_track.get("items", []):
            _insert_paragraph_like(p[94], p[76], item)
    if phases:
        phase_names = [x.get("name", "Fase") for x in phases]
        phase_desc = [x.get("objective", "") for x in phases]
        for index, (hi, body) in enumerate([(86, 87), (88, 89), (90, 91)]):
            if index < len(phases):
                _replace_paragraph(p[hi], phase_names[index]); _replace_paragraph(p[body], phase_desc[index])
            else:
                _remove_paragraph(p[hi]); _remove_paragraph(p[body])
        if len(phases) > 3:
            phase4_heading = p[92].insert_paragraph_before(str(phases[3].get("name", "Produção e estabilização")), style="Heading 3")
            phase4_body = p[92].insert_paragraph_before(str(phases[3].get("objective", "Entrada em produção, acompanhamento e aceite.")))
            phase4_heading.paragraph_format.keep_with_next = True
            phase4_body.paragraph_format.keep_with_next = True
    _replace_paragraph(p[93], "A POPULOS conduzirá o planejamento, a gestão de riscos e dependências e os reportes periódicos, em alinhamento com o ponto focal do cliente.")

    _replace_paragraph(p[96], _doc_text(proposal, "knowledge_transfer", "A passagem de conhecimento abrangerá a arquitetura implantada, as configurações aprovadas, os procedimentos operacionais e a documentação final prevista no escopo."))
    _replace_paragraph(p[98], _doc_text(proposal, "tests", "Os testes validarão decisão de tráfego, failover, políticas de segurança, classificação de automações, risco de conta, desempenho das jornadas e envio de eventos ao SIEM, conforme critérios da seção 10."))
    _set_rows_dynamic(p, range(101, 106), proposal.get("scope", {}).get("excluded", []), p[106])
    _set_rows_dynamic(p, range(110, 118), proposal.get("assumptions", []), p[118])
    _replace_paragraph(p[118], f"8.2 Requisitos — Pessoas ({client})")
    p[108].paragraph_format.page_break_before = False
    client_roles = [
        "Ponto focal executivo: decisões e escalonamentos;",
        "Infraestrutura: acessos a DNS, nuvens e origens;",
        "Segurança/SOC: políticas, integrações e validação das evidências;",
        "Aplicações: apoio aos testes de APIs e jornadas críticas.",
    ]
    _set_rows(p, range(119, 123), client_roles)
    restrictions = proposal.get("restrictions") or [
        "O cronograma depende da disponibilização de acessos, dados de dimensionamento e janelas de mudança;",
        "Mudanças em requisitos, volumetria ou integrações poderão exigir revisão de prazo e escopo;",
        "A proposta não constitui homologação jurídica nem certificação de conformidade regulatória;",
        "A faixa de oito a doze semanas pressupõe acessos no início de cada fase, uma janela de produção aprovada e ausência de alteração material de volumetria.",
    ]
    _set_rows(p, range(124, 128), restrictions)

    _replace_paragraph(p[133], _doc_text(proposal, "schedule_intro", "A implantação está estimada entre oito e doze semanas, contadas a partir da reunião de início, condicionada à liberação de acessos, validação do desenho e disponibilidade das janelas de mudança."))
    _replace_paragraph(p[134], _doc_text(proposal, "schedule_sequence", "A sequência inclui levantamento, desenho, implementação, homologação, entrada em produção e estabilização assistida."))
    _remove_paragraph(p[136])
    _trim_table_columns(t[7], 5)
    headers = ["Fase", "Atividade", "Dependência", "Duração", "Evidência"]
    for c, value in zip(t[7].rows[0].cells, headers): _replace_cell(c, value)
    _repeat_table_header(t[7].rows[0])
    macro = []
    for phase in phases:
        macro.append([phase.get("name", ""), phase.get("objective", ""), phase.get("dependency", "Acessos e validações"), phase.get("duration", ""), ", ".join(phase.get("outputs", []))])
    for wave in optional_phase.get("waves", []):
        components = ", ".join(_product_name(str(item)) for item in wave.get("components", []))
        macro.append([f"{wave.get('name', '')} (opcional)", f"Implantar {components}.", wave.get("dependency", "Contratação da fase"), wave.get("duration", "Definida no Desenho"), wave.get("milestone", "aceite da onda")])
    if fast_track.get("schedule_row"):
        macro.insert(1, list(fast_track["schedule_row"]))
    schedule_rows = list(t[7].rows[1:])
    if len(macro) > len(schedule_rows):
        raise ValueError(f"Cronograma excede a capacidade do template: {len(macro)} fases para {len(schedule_rows)} linhas")
    for row_index, row in enumerate(schedule_rows):
        if row_index < len(macro):
            for c, value in zip(row.cells, macro[row_index]):
                _replace_cell(c, value)
                _set_cell_shading(c, "FFFFFF" if row_index % 2 == 0 else "F3F4F6")
                for paragraph in c.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(31, 42, 94)
                        run.font.size = Pt(7.5)
        else:
            _remove_row(row)

    all_criteria = proposal.get("acceptance_criteria", [])
    criteria = [x.get("criterion", "") for x in all_criteria if x.get("phase") != "optional"]
    reference_criteria = [x.get("criterion", "") for x in all_criteria if x.get("phase") == "optional"]
    _replace_paragraph(p[138], _doc_text(proposal, "acceptance_intro", "A aceitação seguirá os critérios mínimos registrados na matriz de rastreabilidade. A POPULOS propõe, como instrumento complementar de apoio, checklists por entrega, sem prejuízo dos marcos formais de aceite do projeto:"))
    _set_rows_dynamic(p, range(139, 144), criteria, p[144])
    if reference_criteria:
        label = _insert_paragraph_like(p[144], p[138], _doc_text(proposal, "acceptance_optional_intro", "Critérios de referência da fase opcional, aplicáveis somente se contratada:"), copy_run_format=True)
        label.paragraph_format.keep_with_next = True
        for item in reference_criteria:
            _insert_paragraph_like(p[144], p[139], item)
    _replace_paragraph(p[146], _doc_text(proposal, "warranty_intro", "A garantia técnica cobre a correção de defeitos diretamente atribuíveis aos serviços executados pela POPULOS. O prazo consta do quadro Garantia e vigência."))
    _replace_paragraph(p[148], _doc_text(proposal, "closing", "O projeto encerra-se após o aceite e a estabilização assistida. Operação continuada, suporte gerenciado, NOC 24x7 e equipe residente exigem contratação específica."))
    _replace_paragraph(p[149], "Condições de atendimento")
    _replace_paragraph(p[150], "Durante a execução e a garantia, a POPULOS adotará o padrão institucional de atendimento abaixo para falhas atribuíveis aos serviços executados.")

    responsibilities = proposal.get("delivery", {}).get("responsibilities", [])
    p[154].paragraph_format.keep_with_next = True
    role_rows = [list(row) for row in proposal.get("populos_roles", [])] or [
        ["Gerência do PMO", "Supervisão executiva e escalonamento do projeto.", "Sob demanda"],
        ["Executivo de Contas", "Gestão comercial, contratos e governança da conta.", "Sob demanda"],
        ["Gerente de Projeto POPULOS", "Planejamento, coordenação, riscos e reportes.", "Durante o projeto"],
        ["Arquiteto POPULOS", "Desenho e validação da arquitetura Akamai.", "Conforme cronograma"],
        ["Equipe Técnica POPULOS", "Configuração, testes e documentação.", "Conforme cronograma"],
        ["Ponto focal do cliente", "Decisões, acessos e escalonamentos.", "Durante o projeto"],
    ]
    for item in responsibilities:
        role_rows.append([item.get("party", ""), item.get("responsibility", ""), "Conforme plano"])
    responsibility_rows = list(t[11].rows[1:])
    if len(role_rows) > len(responsibility_rows):
        raise ValueError(
            f"Papéis excedem a capacidade do template: {len(role_rows)} itens para {len(responsibility_rows)} linhas"
        )
    for row_index, row in enumerate(responsibility_rows):
        if row_index < len(role_rows):
            for c, value in zip(row.cells, role_rows[row_index]): _replace_cell(c, value)
        else:
            _remove_row(row)

    _replace_paragraph(p[157], f"Todas as informações deste documento são confidenciais e destinam-se exclusivamente à avaliação da proposta por {client} e pela POPULOS.")
    p[156].paragraph_format.page_break_before = False
    p[156].paragraph_format.keep_with_next = True
    p[157].paragraph_format.keep_with_next = True
    _replace_paragraph(p[158], f"Os direitos de propriedade intelectual sobre trabalhos preexistentes, roteiros de instalação, automações e métodos da POPULOS permanecem com a POPULOS. A documentação técnica específica do ambiente será entregue a {client} conforme o escopo aprovado.")
    p[158].paragraph_format.keep_with_next = True
    _prevent_row_split(t[12].rows[0])

    risk_heading = heading2(p[132], "8.4 Riscos e mitigação")
    risk_heading.paragraph_format.space_before = Pt(8)
    risk_table = doc.add_table(rows=1, cols=3)
    risk_heading._p.addnext(risk_table._tbl)
    for c, value in zip(risk_table.rows[0].cells, ["Risco", "Impacto", "Mitigação"]): _replace_cell(c, value)
    for item in proposal.get("risks", []):
        row = risk_table.add_row().cells
        for c, value in zip(row, [item.get("risk", ""), item.get("impact", ""), item.get("mitigation", "")]): _replace_cell(c, value)
    _format_added_table(risk_table, [2.2, 1.6, 2.8])
    _repeat_table_header(risk_table.rows[0])

    options_anchor = trace_heading
    if readiness:
        readiness_heading = heading2(trace_heading, f"3.{readiness_number} Prontidão para eventos críticos")
        readiness_heading.paragraph_format.space_before = Pt(8)
        readiness_heading.paragraph_format.keep_with_next = True
        readiness_intro = _doc_text(proposal, "readiness_intro", "")
        if readiness_intro:
            body_like(trace_heading, readiness_intro)
        for item in readiness:
            _insert_paragraph_like(trace_heading, p[110], item)
        options_anchor = readiness_heading
    if optional_decisions:
        options_heading = heading2(options_anchor, f"3.{options_number} Opções condicionadas")
        options_heading.paragraph_format.space_before = Pt(8)
        options_heading.paragraph_format.keep_with_next = True
        body_like(options_anchor, "Os componentes abaixo não integram a solução principal. Cada um só será incluído se o gatilho indicado for confirmado.")
        for decision in optional_decisions:
            trigger = decision.get("trigger") or decision.get("rationale", "")
            _insert_paragraph_like(options_anchor, p[110], f"{_product_name(str(decision.get('product_id')))} — {trigger}")

    for table_index, heading in ((0, "Aderência e qualificações"), (1, "Modelo de fornecimento"), (2, "Limite da frente"), (5, "Estratégia de migração"), (6, "Cláusulas não aplicáveis"), (8, "Marcos formais de aceite"), (9, "Garantia e vigência")):
        title, body = {
            0: ("Aderência e qualificações", "A equipe será mobilizada com competências compatíveis com DNS, entrega, segurança de aplicações e gestão do projeto; as credenciais nominais serão apresentadas na mobilização."),
            1: ("Modelo de fornecimento", " ".join(filter(None, [_doc_text(proposal, "license_prefix", ""), str(license_supply)]))),
            2: ("Limite da frente", _doc_text(proposal, "callout_limits", "A solução cobre configuração e integração dos componentes recomendados. Alterações no código das aplicações e serviços de terceiros permanecem fora do escopo.")),
            5: ("Estratégia de migração", _doc_text(proposal, "callout_migration", "Implantação em ondas: homologação, produção controlada e estabilização, com janela aprovada, critérios de avanço e plano de reversão.")),
            6: ("Itens que exigem contratação específica", "Operação continuada, NOC 24x7, equipe residente, desenvolvimento de aplicações e serviços gerenciados adicionais não estão incluídos."),
            8: ("Marcos formais de aceite", _doc_text(proposal, "callout_milestones", "Desenho aprovado, homologação concluída, produção validada e documentação final aceita.")),
            9: ("Garantia e vigência", str(warranty)),
        }[table_index]
        _replace_callout(t[table_index].cell(0, 0), title, body)
        _prevent_row_split(t[table_index].rows[0])

    token_replacements = {
        "«CLIENTE»": client,
        "«FABRICANTE»": "Akamai",
        "«CERTIFICAÇÕES»": "qualificações validadas",
        "«X»": "requisito aplicável registrado na matriz de rastreabilidade",
        "«CONTRATANTE ou POPULOS»": "conforme contrato aprovado",
        "«frentes tecnológicas»": ", ".join(front["title"] for front in fronts),
        "«escopo dos testes»": "testes definidos no plano aprovado",
        "«conteúdo»": "conteúdo previsto no escopo",
        "«artefatos de design da frente»": "desenhos e matrizes previstos",
        "«marco de aceite»": "marco formal aprovado",
        "«tipo de entrega»": "entrega",
        "«artefatos»": "evidências e documentação final",
    }
    for paragraph in doc.paragraphs:
        _replace_run_tokens(paragraph, token_replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_run_tokens(paragraph, token_replacements)

    for section in doc.sections:
        for header in (section.header, section.first_page_header):
            for paragraph in header.paragraphs:
                if paragraph.text.strip():
                    _replace_paragraph(paragraph, f"PROPOSTA TÉCNICA | {opportunity} - {client}")
        for footer in (section.footer, section.first_page_footer):
            for paragraph in footer.paragraphs:
                if paragraph.text.strip():
                    _replace_run_tokens(paragraph, {"«FABRICANTE»": "AKAMAI", "Partner|": "Partner |"})

    unresolved: list[str] = []
    all_paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)
    for paragraph in all_paragraphs:
        cleaned = re.sub(
            r"item\s+requisito aplicável registrado na matriz de rastreabilidade(?:\s+do\s+TR)?",
            "requisito aplicável registrado na matriz de rastreabilidade",
            paragraph.text,
            flags=re.IGNORECASE,
        )
        if cleaned != paragraph.text:
            _replace_paragraph(paragraph, cleaned)
    for paragraph in all_paragraphs:
        if "«" in paragraph.text or "[[" in paragraph.text:
            unresolved.append(paragraph.text)
    if unresolved:
        details = "; ".join(dict.fromkeys(unresolved))
        raise ValueError(f"Template contém placeholders não resolvidos: {details}")
    client_text = " ".join(paragraph.text for paragraph in all_paragraphs)
    forbidden = find_vocabulary_violations(client_text, strict=_strict(proposal))
    if forbidden:
        raise ValueError(f"Vocabulário interno ou termo sem tradução no documento: {', '.join(forbidden)}")
    leaked = detect_skill_leakage(client_text)
    if leaked:
        raise ValueError(f"Trechos copiados dos arquivos da skill: {'; '.join(leaked[:5])}")

    for paragraph in doc.paragraphs:
        style_name = str(getattr(paragraph.style, "name", ""))
        if style_name.startswith("Heading") or re.match(r"^\d+(?:\.\d+)*\s", paragraph.text.strip()):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True

    settings = doc.settings._element
    update = settings.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields")
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
    doc.core_properties.title = f"Proposta técnica — {client}"
    doc.core_properties.subject = opportunity
    doc.core_properties.author = "POPULOS Sales Engineering"
    doc.save(output)
    _replace_client_mark(output, client)
    validate_state_against_docx(proposal, output)
    return output
