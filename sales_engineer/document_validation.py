from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable

from docx import Document


_TYPOGRAPHIC_TRANSLATION = str.maketrans({
    "\u00a0": " ",
    "\u00ad": "",
    "\u2018": "'",
    "\u2019": "'",
    "\u201c": '"',
    "\u201d": '"',
    "\u2013": "-",
    "\u2014": "-",
})


def normalize_document_text(value: Any) -> str:
    """Normalize renderer noise without weakening semantic comparisons."""
    text = unicodedata.normalize("NFKC", str(value or "")).translate(_TYPOGRAPHIC_TRANSLATION)
    text = re.sub(r"(?<=\w)-\s*[\r\n]+\s*(?=\w)", "", text)
    text = re.sub(r"(?:^|[\r\n])\s*(?:[•·▪◦]|\d+[.)])\s*", " ", text)
    text = re.sub(r"\s+", " ", text).strip().casefold()
    return text


def extract_docx_text(path: Path) -> str:
    document = Document(path)
    fragments: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            fragments.extend(cell.text for cell in row.cells)
    for section in document.sections:
        for container in (section.header, section.first_page_header, section.footer, section.first_page_footer):
            fragments.extend(paragraph.text for paragraph in container.paragraphs)
    return "\n".join(fragments)


def _client_visible_items(proposal: dict[str, Any]) -> Iterable[tuple[str, str]]:
    scope = proposal.get("scope", {})
    for field in ("included", "deliverables", "excluded"):
        for item in scope.get(field, []):
            if str(item).strip():
                yield f"scope.{field}", str(item)
    for item in proposal.get("acceptance_criteria", []):
        value = item.get("criterion", "") if isinstance(item, dict) else item
        if str(value).strip():
            yield "acceptance_criteria", str(value)
    for item in proposal.get("event_readiness", []):
        text = item.get("text", "") if isinstance(item, dict) else item
        if str(text).strip():
            yield "event_readiness", str(text)
    for wave in (proposal.get("optional_phase") or {}).get("waves", []):
        for field in ("name", "activities", "acceptance"):
            if str(wave.get(field, "")).strip():
                yield f"optional_phase.waves.{field}", str(wave[field])
    for section in proposal.get("sections", []):
        if str(section.get("title", "")).casefold() == "resumo executivo":
            for paragraph in section.get("paragraphs", []):
                yield "sections.resumo_executivo", str(paragraph)
    track = proposal.get("fast_track") or {}
    for text in list(track.get("paragraphs", [])) + list(track.get("items", [])):
        if str(text).strip():
            yield "fast_track", str(text)
    for item in proposal.get("traceability", []):
        if str(item.get("requirement", "")).strip():
            yield "traceability", str(item["requirement"])


SKILL_DIR = Path(__file__).resolve().parents[1] / "skills" / "akamai-proposal-authoring"
LEXICON_PATH = SKILL_DIR / "references" / "lexicon.json"


def load_lexicon(path: Path = LEXICON_PATH) -> dict[str, Any]:
    """Load the single client-language lexicon; a missing file blocks emission."""
    if not path.exists():
        raise ValueError(f"Léxico de linguagem ausente: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def find_vocabulary_violations(text: str, strict: bool, lexicon: dict[str, Any] | None = None) -> list[str]:
    lexicon = lexicon or load_lexicon()
    normalized = normalize_document_text(text)
    terms = list(lexicon.get("forbidden_internal", []))
    if strict:
        terms += lexicon.get("forbidden_internal_strict", [])
        terms += list(lexicon.get("untranslated", {}))
    return [term for term in terms if re.search(rf"\b{re.escape(normalize_document_text(term))}\b", normalized)]


def _shingles(text: str, size: int) -> set[str]:
    words = re.findall(r"\w+", normalize_document_text(text))
    return {" ".join(words[index:index + size]) for index in range(len(words) - size + 1)}


def detect_skill_leakage(text: str, size: int = 8, skill_dir: Path = SKILL_DIR, allowed: Iterable[str] = ()) -> list[str]:
    """Return runs of `size` words copied verbatim from the authoring skill files."""
    skill_text = " ".join(path.read_text(encoding="utf-8") for path in sorted(skill_dir.rglob("*.md")))
    leaked = _shingles(text, size) & _shingles(skill_text, size)
    for exception in allowed:
        leaked -= _shingles(exception, size)
    return sorted(leaked)


def validate_state_against_docx(proposal: dict[str, Any], path: Path) -> None:
    """Block emission when client-visible state disappears from the final DOCX."""
    rendered = normalize_document_text(extract_docx_text(path))
    missing = [
        f"{field}: {value}"
        for field, value in _client_visible_items(proposal)
        if normalize_document_text(value) not in rendered
    ]
    if missing:
        preview = "; ".join(missing[:8])
        suffix = f"; +{len(missing) - 8} item(ns)" if len(missing) > 8 else ""
        raise ValueError(f"Divergência estado × DOCX: {preview}{suffix}")


CASE_RESIDUE_PATH = Path(__file__).resolve().parents[1] / "assets" / "case_residue.json"


def _state_values_text(value: Any) -> str:
    """Every string value of the state (keys excluded), for provenance checks."""
    if isinstance(value, dict):
        return " ".join(_state_values_text(item) for item in value.values())
    if isinstance(value, (list, tuple)):
        return " ".join(_state_values_text(item) for item in value)
    return str(value) if value is not None else ""


def find_case_residue(text: str, proposal: dict[str, Any], path: Path = CASE_RESIDUE_PATH) -> list[str]:
    """Terms known from earlier cases that reach the document without appearing anywhere in this state.

    Such a term can only come from fixed composer text, so it is residue of another case.
    """
    terms = json.loads(path.read_text(encoding="utf-8")).get("terms", [])
    rendered = normalize_document_text(text)
    state_text = normalize_document_text(_state_values_text(proposal))
    found = []
    for term in terms:
        pattern = re.compile(rf"(?<!\w){re.escape(normalize_document_text(term))}(?!\w)")
        if pattern.search(rendered) and not pattern.search(state_text):
            found.append(term)
    return found
