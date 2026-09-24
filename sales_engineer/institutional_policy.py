from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document

from .document_validation import normalize_document_text


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _sha256_text(value: str) -> str:
    return hashlib.sha256(normalize_document_text(value).encode("utf-8")).hexdigest()


def validate_institutional_whitelist(template: Path, manifest_path: Path | None = None) -> dict:
    """Validate institutional exceptions against one exact template version."""
    manifest_path = manifest_path or template.with_name("institutional_whitelist.json")
    if not manifest_path.exists():
        raise ValueError(f"Whitelist institucional ausente: {manifest_path}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    actual_template_hash = _sha256_bytes(template.read_bytes())
    expected_template_hash = str(manifest.get("template_sha256", "")).casefold()
    if actual_template_hash != expected_template_hash:
        raise ValueError(
            "Versão do template não autorizada pela whitelist institucional: "
            f"esperado {expected_template_hash}, obtido {actual_template_hash}"
        )

    document = Document(template)
    sources = {
        "about_populos": " ".join(item.text for item in document.paragraphs[16:19]),
        "confidentiality": " ".join(item.text for item in document.paragraphs[156:159]),
        "sla": " ".join(cell.text for row in document.tables[10].rows for cell in row.cells),
    }
    for block_id, expected_hash in manifest.get("blocks", {}).items():
        if block_id not in sources:
            raise ValueError(f"Bloco institucional desconhecido: {block_id}")
        actual_hash = _sha256_text(sources[block_id])
        if actual_hash != str(expected_hash).casefold():
            raise ValueError(
                f"Bloco institucional alterado sem nova versão: {block_id} "
                f"(esperado {expected_hash}, obtido {actual_hash})"
            )
    return manifest
